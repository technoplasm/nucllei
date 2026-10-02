"""sandbox jail の隔離テスト (権限昇格 / デバイス・ハードウェア・システムの見え方)。
加えて、jail の制約 (rootfs ro・HOME=/workspace) に引っかかりやすい同梱ツール
(LibreOffice headless 等) が jail の中で実際に動くことも見る。

これは稼働中の privileged コンテナに対する統合テストで、service が組み立てる nsjail の設定
(sandbox/README.md「セキュリティモデル」) を jail の *中から* 実際に突いて確認する。
ユニットテストではない。

到達経路: service のポートは host 非公開なので、コンテナ内の Python で
localhost:8000/exec を叩く (`docker exec` 越し)。追加依存なし (docker CLI + stdlib)。
コンテナが無い / 落ちていれば全テスト skip。

コンテナ名は配備側の compose が付けるので env で差し替えられる (既定は現行配備の
`llens-sandbox`)。

実行:
    uv run task sandbox-test
    python3 sandbox/tests/test_isolation.py     # pytest 不要 (末尾ランナー)
    SANDBOX_CONTAINER=<name> uv run task sandbox-test
"""

import json
import os
import subprocess

CONTAINER = os.environ.get("SANDBOX_CONTAINER", "llens-sandbox")
USER_ID = "isolationtest"  # _USER_ID_RE に合う専用テストユーザー

# コンテナ内で /exec を叩く固定クライアント。ユーザーコードは stdin の JSON に載せ、
# この -c 文字列自体は不変にして quoting 事故を避ける。
# 共有トークン (#132) は**コンテナ自身の env から取る** — 配備側がサービスに渡した値が
# そのまま正解なので、テスト側に値を持たせない (ずれようがない)。認証は opt-in なので
# 未設定のコンテナでは header の中身が空になるが、その場合は service 側も認証を掛けない
# ためそのまま通る。
_CLIENT = (
    "import os,sys,urllib.request\n"
    "d=sys.stdin.read()\n"
    "r=urllib.request.urlopen(urllib.request.Request("
    "'http://localhost:8000/exec',data=d.encode(),"
    "headers={'Content-Type':'application/json',"
    "'Authorization':'Bearer '+os.environ.get('SANDBOX_TOKEN','')}),timeout=90)\n"
    "sys.stdout.write(r.read().decode())\n"
)


def run(lang: str, code: str, timeout_s: int = 30) -> dict:
    """jail 内で code を実行し /exec の JSON を返す。lang は "python" | "shell"。"""
    payload = json.dumps(
        {"user_id": USER_ID, "lang": lang, "code": code, "timeout_s": timeout_s}
    )
    p = subprocess.run(
        ["docker", "exec", "-i", CONTAINER, "/opt/app/bin/python", "-c", _CLIENT],
        input=payload,
        capture_output=True,
        text=True,
        timeout=timeout_s + 60,
    )
    if p.returncode != 0:
        raise RuntimeError(
            f"docker exec {CONTAINER} rc={p.returncode}: {p.stderr[:300]}"
        )
    return json.loads(p.stdout)


def out(lang: str, code: str) -> str:
    """成功実行の stdout を返す (exit_code!=0 なら stderr 付きで落とす)。"""
    r = run(lang, code)
    assert r["exit_code"] == 0, f"exit={r['exit_code']} stderr={r['stderr'][:300]!r}"
    return r["stdout"]


def _service_reachable() -> bool:
    try:
        return run("python", "print('ok')")["stdout"].strip() == "ok"
    except Exception:
        return False


AVAILABLE = _service_reachable()
SKIP_REASON = f"sandbox container {CONTAINER!r} not reachable"

try:
    import pytest

    pytestmark = pytest.mark.skipif(not AVAILABLE, reason=SKIP_REASON)
except ImportError:  # __main__ ランナーだけで回すとき
    pytest = None


# ---------------------------------------------------------------------------
# 権限昇格
# ---------------------------------------------------------------------------
def test_runs_as_unprivileged_nobody():
    # jail 内 uid/gid は全ユーザー共通の 65534 (nobody)。root では動かない。
    assert out("python", "import os;print(os.getuid(),os.getgid())").split() == [
        "65534",
        "65534",
    ]


def test_cannot_setuid_root():
    # root へ上がれない (CAP_SETUID を持たない)。setuid(0) は PermissionError。
    assert (
        out(
            "python",
            "import os\ntry:\n os.setuid(0);print('ESCALATED')\n"
            "except PermissionError:\n print('DENIED')",
        ).strip()
        == "DENIED"
    )


def test_system_dirs_read_only():
    # rootfs (/usr /etc /opt/venv) は read-only bind。書けるのは /workspace と /tmp のみ。
    r = run(
        "shell",
        "for d in /usr /etc /opt/venv; do touch $d/pwned 2>&1; done; "
        "ls /usr/pwned /etc/pwned /opt/venv/pwned 2>&1",
    )
    assert r["stdout"].count("Read-only file system") == 3
    assert "No such file" in r["stdout"]  # どのファイルも作られていない


def test_ro_mount_root_is_read_only():
    # 参照用データの置き場 (/mnt、#138): 入れたものは jail から読めるが**書けない**。
    # 何も入れていない配備 (= 既定) では確かめる対象が無いので何もしない。
    listed = run("shell", "ls /mnt 2>/dev/null")["stdout"].split()
    if not listed:
        print("  note: /mnt に何も無い配備なので確認をスキップ (volume 未設定 = 既定)")
        return
    target = f"/mnt/{listed[0]}"
    # touch は dir/file どちらでも書き込み扱い (utimensat) なので RO なら EROFS になる
    r = run("shell", f"[ -e {target} ] && echo VISIBLE; touch {target} 2>&1")
    assert "VISIBLE" in r["stdout"], r["stdout"]
    assert "Read-only file system" in r["stdout"], r["stdout"]


def test_no_host_filesystem_access():
    # 他ユーザーの workspace / イメージ置き場 (ホスト /workspaces) は jail に bind されない。
    r = run("shell", "ls -d /workspaces /workspaces/.images 2>&1")
    assert r["exit_code"] != 0
    assert "No such file or directory" in r["stdout"]


# ---------------------------------------------------------------------------
# デバイス / ハードウェア / システムの見え方
# ---------------------------------------------------------------------------
def test_dev_contains_only_whitelisted_nodes():
    # jail の / は新規 tmpfs。/dev は DEV_BINDS の 5 ノードだけを bind する。
    assert sorted(out("shell", "ls /dev").split()) == [
        "full",
        "null",
        "random",
        "urandom",
        "zero",
    ]


def test_host_block_and_gpu_devices_absent():
    # コンテナは privileged でホストのデバイスを見られるが、jail の /dev はそれを剥がす。
    # ディスク実体・GPU・カーネルメモリのノードは一切見えてはならない。
    probes = [
        "/dev/sda",
        "/dev/loop0",
        "/dev/nvme0n1",
        "/dev/mem",
        "/dev/kmsg",
        "/dev/nvidia0",
        "/dev/dm-0",
    ]
    script = "".join(f"[ -e {d} ] && echo PRESENT {d}\n" for d in probes) + "echo done"
    r = out("shell", script)
    assert "PRESENT" not in r


def test_network_is_unreachable():
    # network namespace は空 (lo も上げない)。外部への到達は構造的に不可。
    r = out(
        "python",
        "import socket\ntry:\n"
        " s=socket.socket();s.settimeout(3);s.connect(('1.1.1.1',80));print('CONNECTED')\n"
        "except OSError as e:\n print('BLOCKED', e.errno)",
    )
    assert r.startswith("BLOCKED")


def test_disk_shows_quota_not_host():
    # df は per-user quota (既定 10GB) を返す。ホストの実 1.8TB は見えない
    # (「なんでも書ける」誤認防止)。ルート tmpfs は 16MB 固定。
    ws_total = int(out("shell", "df -B1 --output=size /workspace | tail -1").strip())
    assert 1 * 1024**3 <= ws_total <= 64 * 1024**3, ws_total  # GB オーダー、TB ではない
    root_total = int(out("shell", "df -B1 --output=size / | tail -1").strip())
    assert root_total == 16 * 1024 * 1024


def test_pid_namespace_hides_host_processes():
    # PID namespace が独立。jail 内からはホスト (uvicorn/nsjail 等) のプロセスが見えず、
    # 自プロセスは pid 1 付近。/proc に並ぶ pid はごく少数。
    line = out(
        "python",
        "import os;pids=[p for p in os.listdir('/proc') if p.isdigit()];"
        "print(os.getpid(),len(pids))",
    )
    my_pid, visible = (int(x) for x in line.split())
    assert my_pid <= 3
    assert visible <= 5


# ---------------------------------------------------------------------------
# 同梱ツールが jail の制約 (rootfs ro・HOME=/workspace・無ネットワーク) の下で動く
# ---------------------------------------------------------------------------
# LibreOffice は起動のたびにプロファイルを書くので、rootfs ro の jail では置き場を誤ると
# 死ぬか、HOME (= 永続 workspace) を汚す。/usr/local/bin/soffice の wrapper が /tmp へ
# 向けている (sandbox/soffice-wrapper.sh)。ここではそれを外から確かめる。
_OFFICE_FIXTURE = (
    "set -e; mkdir -p /tmp/o && cd /tmp/o\n"
    "python - <<'P'\n"
    "import docx, openpyxl\n"
    "d = docx.Document()\n"
    "for i in range(20): d.add_paragraph(f'退院時サマリー 第{i}段落 discharge summary {i}')\n"
    "d.save('t.docx')\n"
    "wb = openpyxl.Workbook(); ws = wb.active; ws.title = '検査'\n"
    "ws.append(['項目', '値']); ws.append(['WBC', 5400]); wb.save('t.xlsx')\n"
    "P\n"
    # 97-2003 形式は LibreOffice 自身に作らせる (= 書き出しも jail で通ることの確認を兼ねる)
    "soffice --headless --convert-to doc t.docx >/dev/null 2>&1\n"
    "soffice --headless --convert-to xls t.xlsx >/dev/null 2>&1\n"
)


def test_soffice_converts_legacy_office_in_jail():
    # .doc/.xls を PDF へ。日本語が Noto CJK で埋め込まれて落ちないこと、
    # そして workspace にプロファイル類 (.config / .local 等の dotfile) を残さないこと。
    ws_before = out("shell", "ls -A /workspace")
    r = run(
        "shell",
        _OFFICE_FIXTURE + "file -b t.doc t.xls | cut -c1-30\n"
        # 出力名は入力の basename なので、t.doc と t.xls は outdir を分ける (同じだと上書き)
        + "soffice --headless --convert-to pdf --outdir d t.doc >/dev/null 2>&1\n"
        + "soffice --headless --convert-to pdf --outdir x t.xls >/dev/null 2>&1\n"
        + "pdftotext x/t.pdf - | grep -m1 WBC\n"
        + "pdftotext d/t.pdf - | head -1\n"
        # python-docx 既定の和文指定は明朝系なので IPA 明朝に落ちる (IPA が無い image では Noto)
        + "pdffonts d/t.pdf | grep -c -E 'NotoSansCJK|IPA'\n",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    lines = r["stdout"].splitlines()
    # 1・2 行目: 本当に OLE2 (97-2003 形式) になっている
    assert lines[0].startswith("Composite Document File V2"), lines
    assert lines[1].startswith("Composite Document File V2"), lines
    assert lines[2].startswith("WBC"), lines
    assert "退院時サマリー" in lines[3], lines
    assert int(lines[-1]) >= 1  # 日本語フォントが埋め込まれている
    assert out("shell", "ls -A /workspace") == ws_before  # HOME を汚していない


def test_soffice_maps_ms_fonts_to_ipa_in_jail():
    # 「ＭＳ 明朝 / ＭＳ ゴシック」指定 (全角・半角表記とも) の古い Word を PDF 化したとき、
    # fontconfig の alias (Dockerfile の 64-nucllei-fonts.conf) で字幅の合う IPA に寄ること。
    # alias が無いと LibreOffice は DejaVu と IPA を混ぜ、P (プロポーショナル) も区別されない
    r = run(
        "shell",
        "set -e; mkdir -p /tmp/f && cd /tmp/f\n"
        "python - <<'P'\n"
        "import docx\n"
        "from docx.oxml.ns import qn\n"
        "d = docx.Document()\n"
        "for f in ['ＭＳ 明朝', 'MS Mincho', 'ＭＳ ゴシック', 'ＭＳ Ｐゴシック']:\n"
        "    r = d.add_paragraph().add_run(f'診療情報提供書 {f}'); r.font.name = f\n"
        "    r._element.rPr.rFonts.set(qn('w:eastAsia'), f)\n"
        "d.save('ms.docx')\n"
        "P\n"
        "soffice --headless --convert-to pdf ms.docx >/dev/null 2>&1\n"
        "pdffonts ms.pdf | tail -n +3 | awk '{print $1}' | cut -d+ -f2 | sort\n",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    assert r["stdout"].split() == ["IPAGothic", "IPAMincho", "IPAPGothic"], r["stdout"]


def test_legacy_office_text_extractors_in_jail():
    # 軽量抽出 (soffice を起動しない経路): antiword (.doc) / xls2csv (.xls) / xlrd / olefile
    r = run(
        "shell",
        _OFFICE_FIXTURE
        + "antiword -m UTF-8.txt t.doc | grep -m1 退院時\n"
        + "xls2csv -d utf-8 t.xls | grep WBC\n"
        + "python -c \"import xlrd,olefile;print(olefile.isOleFile('t.doc'),"
        + "xlrd.open_workbook('t.xls').sheet_names())\"\n",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    assert "退院時サマリー 第0段落" in r["stdout"]
    assert '"WBC","5400"' in r["stdout"]
    assert "True ['検査']" in r["stdout"]


# ---------------------------------------------------------------------------
# レガシーデータ層 (Dockerfile の apt 層 + requirements.txt の「レガシーデータ」)
# ---------------------------------------------------------------------------
# LZH は lhasa が展開専用なので、無圧縮 (-lh0-) の level-0 書庫を手で組む。名前は
# Windows 産の実物と同じく Shift_JIS で入れる (lha / 7z は生バイトのまま出して化けるが、
# unar / lsar は推定して UTF-8 に直す — README の使い分けの根拠)
_LEGACY_FIXTURE = (
    "python - <<'P'\n"
    "import struct\n"
    "def crc16(b):\n"
    "    c = 0\n"
    "    for x in b:\n"
    "        c ^= x\n"
    "        for _ in range(8): c = (c >> 1) ^ 0xA001 if c & 1 else c >> 1\n"
    "    return c\n"
    "def entry(name, data):\n"
    "    n = name.encode('cp932')\n"
    "    h = b'-lh0-' + struct.pack('<IIIBB', len(data), len(data), 0x5A210000, 0x20, 0)\n"
    "    h += bytes([len(n)]) + n + struct.pack('<H', crc16(data))\n"
    "    return bytes([len(h), sum(h) & 0xFF]) + h + data\n"
    "open('t.lzh', 'wb').write(entry('検査結果.txt', b'WBC 5400\\r\\n') + b'\\0')\n"
    "import io, zipfile\n"
    "b = io.BytesIO(); z = zipfile.ZipFile(b, 'w'); z.writestr('AAAAAAAA.txt', 'Hb 13.2'); z.close()\n"
    "open('s.zip', 'wb').write(b.getvalue().replace(b'AAAAAAAA', '検査結果'.encode('cp932')))\n"
    "open('t.rtf', 'w').write(r\"{\\rtf1\\ansi\\ansicpg932{\\fonttbl{\\f0\\fcharset128 MS Mincho;}}\\f0 \\'8c\\'9f\\'8d\\'b8 WBC\\par}\")\n"
    "P\n"
)


def test_legacy_data_shell_tools_in_jail():
    r = run(
        "shell",
        "set -e; mkdir -p /tmp/l && cd /tmp/l\n"
        + _LEGACY_FIXTURE
        + "printf 'MZ\\0' | xxd | cut -c1-20\n"
        + "printf 'MZ' | hexdump -C | head -1\n"
        + "mkdir a && (cd a && lha xq ../t.lzh) && cat a/*\n"
        + "lsar t.lzh | tail -1\n"
        + "unar -q -o u s.zip && ls u\n"
        # RTF (cp932): unrtf --text は和文を落とすので --html を pandoc で平文に
        + "unrtf --html t.rtf 2>/dev/null | pandoc -f html -t plain\n"
        + "sqlite3 x.db \"create table a(k,v); insert into a values('WBC',5400)\" "
        + "&& sqlite3 -csv x.db 'select * from a'\n"
        + "echo '<a><b>検査</b></a>' | xmllint --xpath '//b/text()' -\n"
        + "printf 'a\\r\\n' > c.txt && dos2unix -q c.txt && file -b c.txt\n"
        + "7z a -bd -ppw z.7z c.txt >/dev/null && 7z x -bd -ppw -oz z.7z >/dev/null && cat z/c.txt\n"
        + "python -c \"from fpdf import FPDF;p=FPDF();p.add_page();p.add_page();p.output('p.pdf')\"\n"
        + "qpdf --encrypt u o 256 -- p.pdf e.pdf && qpdf --password=u --decrypt e.pdf d.pdf "
        + "&& qpdf --show-npages d.pdf\n"
        # Access / cab は手元で作れないので起動確認だけ (実ファイルでの確認は README に記録)
        + "mdb-tables --version; mdb-export --version | cut -c1-8; mdb-schema --version | cut -c1-8\n"
        + "cabextract --version | cut -c1-10\n",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    lines = [ln.rstrip() for ln in r["stdout"].splitlines() if ln.strip()]
    assert lines == [
        "00000000: 4d5a 00",
        "00000000  4d 5a                                             |MZ|",
        "WBC 5400",
        "検査結果.txt",
        "検査結果.txt",
        "検査 WBC",  # unrtf --html は fonttbl の \\fcharset128 で cp932 と解く
        "WBC,5400",
        "検査",
        "ASCII text",
        "a",
        "2",
        "mdbtools v1.0.0",
        "mdbtools",
        "mdbtools",
        "cabextract",
    ], r["stdout"]


def test_legacy_data_python_libs_in_jail():
    # 読むだけのもの (dbf / sav / pdf の表) は venv 内の書き手で作って戻す
    r = run(
        "python",
        "import datetime, io, os, struct\n" "os.chdir('/tmp')\n"
        # dBASE III (.dbf)。言語ドライバ 0x13 = cp932
        "rec = b' ' + '白血球'.encode('cp932').ljust(10) + b'5400'.ljust(6)\n"
        "hdr = struct.pack('<BBBBIHH20x', 3, 124, 10, 2, 1, 32 + 64 + 1, 17)\n"
        "hdr = hdr[:29] + b'\\x13' + hdr[30:]\n"
        "for n, t, l in [(b'NAME', b'C', 10), (b'VAL', b'N', 6)]:\n"
        "    hdr += struct.pack('<11sc4xBB14x', n, t, l, 0)\n"
        "open('k.dbf', 'wb').write(hdr + b'\\r' + rec + b'\\x1a')\n"
        "import dbfread\n"
        "print([dict(r) for r in dbfread.DBF('k.dbf', encoding='cp932')])\n"
        "import hl7\n"
        "m = hl7.parse('MSH|^~\\\\&|HIS|H|LAB|H|20261002||ORU^R01|1|P|2.5\\r"
        "PID|||123^^^H||北大^太郎\\rOBX|1|NM|WBC||5400|/uL')\n"
        "print(m.segment('PID')[5], m.segment('OBX')[5])\n"
        "import pandas as pd, pyreadstat\n"
        "pyreadstat.write_sav(pd.DataFrame({'wbc': [5400.0]}), 'a.sav', column_labels=['白血球数'])\n"
        "df, meta = pyreadstat.read_sav('a.sav'); print(df.wbc.tolist(), meta.column_labels)\n"
        "from fpdf import FPDF\n"
        "p = FPDF(); p.add_page(); p.set_font('helvetica', size=10)\n"
        "with p.table() as t:\n"
        "    for row in [('item', 'value'), ('WBC', '5400')]:\n"
        "        r = t.row(); [r.cell(c) for c in row]\n"
        "p.output('t.pdf')\n"
        "import pdfplumber\n"
        "with pdfplumber.open('t.pdf') as pdf: print(pdf.pages[0].extract_table())\n"
        "import xlwt, xlrd\n"
        "wb = xlwt.Workbook(); wb.add_sheet('検査').write(0, 0, 'WBC'); wb.save('w.xls')\n"
        "print(xlrd.open_workbook('w.xls').sheet_names())\n"
        "import mojimoji, jpholiday, extract_msg\n"
        "print(mojimoji.zen_to_han('ＷＢＣ５４００'), mojimoji.han_to_zen('ｹﾝｻ'))\n"
        "print(jpholiday.is_holiday_name(datetime.date(2026, 11, 3)))\n"
        "print(callable(extract_msg.openMsg))\n",
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    assert r["stdout"].splitlines() == [
        "[{'NAME': '白血球', 'VAL': 5400}]",
        "北大^太郎 5400",
        "[5400.0] ['白血球数']",
        "[['item', 'value'], ['WBC', '5400']]",
        "['検査']",
        "WBC5400 ケンサ",
        "文化の日",
        "True",
    ], r["stdout"]


# ---------------------------------------------------------------------------
# OCR・画像・研究データ層 (Dockerfile の apt 層 + requirements.txt の「OCR・画像」以下)
# ---------------------------------------------------------------------------
# スキャン風の素材: Noto で和文を描いた A4 (150dpi) の画像 3 枚 → img2pdf で画像だけの PDF
_SCAN_FIXTURE = """set -e; mkdir -p /tmp/s && cd /tmp/s
python - <<'P'
from PIL import Image, ImageDraw, ImageFont
f = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 48)
for i in range(3):
    im = Image.new('L', (1240, 1754), 255)
    d = ImageDraw.Draw(im)
    for j, t in enumerate(['血圧 120/80 mmHg', f'診療情報提供書 {i + 1}ページ']):
        d.text((100, 150 + j * 120), t, font=f, fill=0)
    im.save(f'p{i}.png', dpi=(150, 150))
P
img2pdf p0.png p1.png p2.png -o scan.pdf
"""


def test_ocr_tesseract_and_ocrmypdf_in_jail():
    # tesseract (jpn) が画像を読み、ocrmypdf が画像だけの PDF を検索可能 PDF にする。
    # ocrmypdf は一時ファイルを TMPDIR (= jail の揮発 /tmp) に置く — workspace を汚さないこと
    ws_before = out("shell", "ls -A /workspace")
    r = run(
        "shell",
        _SCAN_FIXTURE + """tesseract p0.png - -l jpn 2>/dev/null | grep -m1 血圧
pdftotext scan.pdf - | tr -d '\\f\\n' | wc -c
ocrmypdf -q -l jpn --force-ocr scan.pdf ocr.pdf
pdftotext ocr.pdf - | grep -c 診療情報提供書
""",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    lines = r["stdout"].split("\n")
    assert lines[0].replace(" ", "") == "血圧120/80mmHg", lines
    assert lines[1] == "0", lines  # OCR 前はテキスト層が無い
    assert lines[2] == "3", lines  # 3 ページとも読めている
    assert out("shell", "ls -A /workspace") == ws_before


def test_image_and_doc_shell_tools_in_jail():
    r = run(
        "shell",
        """set -e; mkdir -p /tmp/i && cd /tmp/i
python - <<'P'
import docx, numpy as np, segno
from PIL import Image
segno.make('患者ID:12345', micro=False).save('q.png', scale=8)
# 300dpi の A4 (img2pdf は再圧縮しない) を gs /ebook (150dpi) で縮める
a = np.random.default_rng(0).integers(0, 255, (1754, 1240), np.uint8)
Image.fromarray(a).resize((2480, 3508)).save('hi.png', dpi=(300, 300))
Image.fromarray(a[:800, :800]).convert('RGB').save('n.jpg')
d = docx.Document(); d.core_properties.author = '北大 太郎'; d.add_paragraph('x'); d.save('m.docx')
P
img2pdf hi.png -o big.pdf
zbarimg -q q.png
exiftool -s3 -Creator m.docx
exiftool -s3 -ImageWidth n.jpg
gs -q -sDEVICE=pdfwrite -dPDFSETTINGS=/ebook -dNOPAUSE -dBATCH -sOutputFile=small.pdf big.pdf
test $(stat -c %s small.pdf) -lt $(( $(stat -c %s big.pdf) / 2 )) && echo shrunk
echo 'digraph{a[label="外来"];b[label="入院"];a->b}' | dot -Tpdf -o g.pdf
pdffonts g.pdf | tail -n +3 | awk '{print $1}' | cut -d+ -f2
echo 'digraph{a[label="外来"]}' | dot -Tpng -o g.png && file -b g.png | cut -d, -f1
fc-match -f '%{family[0]}\\n' 'sans-serif:lang=ja'
mkdir r && cd r && git init -q && echo 血圧は高い > a.txt && git add a.txt && git commit -qm 1
echo 血圧は低い > a.txt && git diff --word-diff | tail -1 && git log --format='%an' && cd ..
echo a b c > w1 && echo a x c > w2 && { wdiff w1 w2 || true; }
tree r | tail -1
tiffinfo 2>&1 | head -1 | cut -c1-7
mediainfo --Version | tail -1 | cut -c1-12
csvlook --version | cut -c1-7
""",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    assert r["stdout"].splitlines() == [
        "QR-Code:患者ID:12345",
        "北大 太郎",  # docx のメタ (zip 内) は libarchive-zip-perl が要る
        "800",
        "shrunk",
        # dot の既定 (Times) → serif の和文は Noto Serif CJK のまま (IPA は MS 指定のときだけ)
        "NotoSerifCJKjp-Regular",
        "PNG image data",
        "Noto Sans CJK JP",
        "[-血圧は高い-]{+血圧は低い+}",
        "sandbox",  # /etc/gitconfig の既定 identity (jail の nobody には名前もメールも無い)
        "a [-b-] {+x+} c",
        "1 directory, 1 file",
        "LIBTIFF",
        "MediaInfoLib",
        "csvlook",
    ], r["stdout"]
    # git が HOME (= 永続 workspace) に .gitconfig を書いていない
    assert ".gitconfig" not in out("shell", "ls -A /workspace")


def test_research_python_libs_in_jail():
    r = run(
        "python",
        """import io, os
os.chdir('/tmp')
import numpy as np, cv2
im = np.full((60, 200), 200, np.uint8)
cv2.putText(im, 'WBC', (10, 45), cv2.FONT_HERSHEY_SIMPLEX, 1.5, 30, 3)
cv2.imwrite('c.png', im)
_, b = cv2.threshold(cv2.imread('c.png', cv2.IMREAD_GRAYSCALE), 0, 255,
                     cv2.THRESH_BINARY + cv2.THRESH_OTSU)
print(sorted(set(b.ravel().tolist())))
import pandas as pd
from tableone import TableOne
df = pd.DataFrame({'age': [60, 70, 65, 80, 55, 72], 'sex': list('MFMFMF'), 'g': list('AAABBB')})
print(TableOne(df, columns=['age', 'sex'], categorical=['sex'], groupby='g').tableone.shape[0])
import pingouin as pg, scikit_posthocs as sp
print(round(float(pg.ttest([1, 2, 3, 4, 5], [3, 4, 5, 6, 7])['p_val'].iloc[0]), 3))
print(sp.posthoc_dunn([[1, 2, 3], [4, 5, 6], [7, 8, 9]]).shape)
import neologdn, pykakasi
print(neologdn.normalize('ｹﾝｻ　　ｹｯｶ'), ''.join(x['hira'] for x in pykakasi.kakasi().convert('血圧')))
import h5py, nibabel as nib
with h5py.File('a.h5', 'w') as f: f['x'] = np.arange(3)
with h5py.File('a.h5') as f: print(f['x'][:].tolist())
nib.save(nib.Nifti1Image(np.zeros((2, 2, 2), np.int16), np.eye(4)), 'a.nii.gz')
print(nib.load('a.nii.gz').shape)
import img2pdf, pdf2image, pytesseract, pyzbar.pyzbar, segno
from PIL import Image
segno.make('WBC5400', micro=False).save('q.png', scale=8)
pages = pdf2image.convert_from_bytes(img2pdf.convert('q.png'), dpi=72)
print(len(pages), [d.data.decode() for d in pyzbar.pyzbar.decode(Image.open('q.png'))])
print(pytesseract.get_tesseract_version().major >= 5)
import graphviz, yaml, xmltodict, docxtpl, ocrmypdf
print(graphviz.Digraph(body=['a -> b']).pipe('svg')[:5], yaml.safe_load('a: 1'), xmltodict.parse('<a>1</a>'))
""",
        timeout_s=120,
    )
    assert r["exit_code"] == 0, r["stderr"][-500:]
    assert r["stdout"].splitlines() == [
        "[0, 255]",
        "4",
        "0.081",
        "(3, 3)",
        "ケンサケッカ けつあつ",  # neologdn は和文間の空白を詰める
        "[0, 1, 2]",
        "(2, 2, 2)",
        "1 ['WBC5400']",
        "True",
        "b'<?xml' {'a': 1} {'a': '1'}",
    ], r["stdout"]


if __name__ == "__main__":
    if not AVAILABLE:
        print(f"SKIP (all): {SKIP_REASON}")
        raise SystemExit(0)
    tests = [
        v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)
    ]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"PASS {t.__name__}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"FAIL {t.__name__}: {e!r}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    raise SystemExit(1 if failed else 0)
